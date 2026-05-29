#!/usr/bin/env python3
"""Generate internship report (.docx) for CMH Career Hub stage project."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent / "Rapport_Stage_AYA_EL_OUAHABI_CMH_CareerHub.docx"

TEAM = [
    "AYA EL OUAHABI",
    "AYA AOULAD AHMIDOU",
    "ZAHIRA KTAIB",
    "IHSSAN YAKOUBI",
]


def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    style.paragraph_format.space_after = Pt(6)


def add_simple_footer(doc: Document) -> None:
    """Pied de page texte simple (compatible Word/LibreOffice)."""
    section = doc.sections[0]
    footer = section.footer
    for para in list(footer.paragraphs):
        p_element = para._element
        p_element.getparent().remove(p_element)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CMCTTA - Rapport de stage - CMH Career Hub - AYA EL OUAHABI")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def add_centered(doc: Document, text: str, *, bold=False, size=12, space_after=6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_paragraph(doc: Document, text: str, *, bold=False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths=None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, cell in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = cell
            for p in table.rows[r_idx].cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)
    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = width


def build_cover_page(doc: Document) -> None:
    add_centered(doc, "[Logo OFPPT]", bold=True, size=14, space_after=24)
    add_centered(doc, "Établissement : Centre de Métiers et de Compétences TTA (CMCTTA)", space_after=6)
    add_centered(doc, "Filière : Développement Digital", space_after=6)
    add_centered(doc, "Option : Web Full Stack", space_after=24)
    add_centered(doc, "RAPPORT DE STAGE DE FIN DE FORMATION", bold=True, size=16, space_after=24)
    add_centered(doc, "Stagiaire : AYA EL OUAHABI", bold=True, space_after=12)
    add_centered(doc, "Équipe projet :", bold=True, space_after=6)
    for member in TEAM:
        add_centered(doc, f"• {member}", space_after=4)
    add_centered(doc, "", space_after=12)
    add_centered(doc, "Entreprise d'accueil : Cloud Marketing Hub (CMH) — ....", space_after=6)
    add_centered(doc, "Période du stage : du .... au ....", space_after=6)
    add_centered(doc, "Encadrant pédagogique : ....", space_after=6)
    add_centered(doc, "Tuteur professionnel : ....", space_after=6)
    add_centered(doc, "Année de formation : 2025–2026", space_after=6)
    doc.add_page_break()


def build_remerciements(doc: Document) -> None:
    add_heading(doc, "REMERCIEMENTS", 1)
    paragraphs = [
        "Au terme de cette période de stage, je tiens à exprimer ma profonde gratitude à toutes les personnes qui ont contribué à la réussite de cette expérience professionnelle.",
        "Je remercie sincèrement Cloud Marketing Hub (CMH) pour m'avoir accueillie au sein de son équipe et pour m'avoir confié la réalisation du projet CMH Career Hub, une plateforme web de gestion des promotions de formation et du suivi des candidats. Cette opportunité m'a permis de mettre en pratique les compétences acquises au CMCTTA dans un contexte professionnel concret.",
        "Mes remerciements s'adressent également à mon tuteur professionnel, ...., pour son encadrement, ses conseils techniques et sa disponibilité tout au long du stage. Son accompagnement m'a aidée à progresser rapidement et à adopter les bonnes pratiques du développement web en entreprise.",
        "Je remercie aussi l'équipe pédagogique de l'OFPPT au CMCTTA, en particulier mon encadrant ...., pour le suivi pédagogique, les orientations méthodologiques et le soutien apporté avant et pendant le stage.",
        "Enfin, je remercie mes collègues de stage — AYA AOULAD AHMIDOU, ZAHIRA KTAIB et IHSSAN YAKOUBI — pour la collaboration, l'entraide et le travail d'équipe qui ont permis de mener à bien ce projet ambitieux.",
    ]
    for text in paragraphs:
        add_paragraph(doc, text)
    doc.add_page_break()


def build_sommaire(doc: Document) -> None:
    add_heading(doc, "SOMMAIRE", 1)
    rows = [
        ["1", "Page de garde", "i"],
        ["2", "Remerciements", "ii"],
        ["3", "Sommaire", "iii"],
        ["4", "Introduction Générale", "1"],
        ["5", "Présentation de l'entreprise", "3"],
        ["6", "Présentation du service d'accueil", "5"],
        ["7", "Missions et tâches réalisées", "7"],
        ["8", "Outils, technologies et méthodes utilisés", "15"],
        ["9", "Compétences acquises", "18"],
        ["10", "Difficultés rencontrées et solutions apportées", "21"],
        ["11", "Bilan personnel", "24"],
        ["12", "Conclusion Générale", "26"],
        ["13", "Annexes (UML, Planning, Captures)", "28"],
        ["14", "Bibliographie / Webographie", "32"],
    ]
    add_table(doc, ["N°", "Partie", "Page"], rows, [Cm(1.5), Cm(12), Cm(2.5)])
    doc.add_page_break()


def build_introduction(doc: Document) -> None:
    add_heading(doc, "1. INTRODUCTION GÉNÉRALE", 1)
    add_heading(doc, "1.1 Contexte de la formation", 2)
    add_paragraph(
        doc,
        "La filière Développement Digital, option Web Full Stack, proposée par l'OFPPT au CMCTTA, vise à former des techniciens spécialisés capables de concevoir, développer et maintenir des applications web complètes. Cette formation couvre l'ensemble de la chaîne de développement : interfaces utilisateur (front-end), logique métier et API (back-end), bases de données, sécurité, et bonnes pratiques de développement.",
    )
    add_paragraph(
        doc,
        "Au cours de la formation, nous avons étudié les langages web fondamentaux (HTML, CSS, JavaScript), les frameworks modernes (React, Laravel), la gestion de bases de données relationnelles (MySQL), ainsi que les méthodologies de travail en équipe et les outils de versioning (Git/GitHub).",
    )
    add_heading(doc, "1.2 Importance du stage", 2)
    add_bullets(
        doc,
        [
            "Appliquer les connaissances théoriques à un projet réel ;",
            "Découvrir l'environnement professionnel et ses exigences ;",
            "Développer des compétences techniques et comportementales ;",
            "Préparer l'insertion sur le marché du travail.",
        ],
    )
    add_paragraph(
        doc,
        "Notre stage s'est déroulé au sein de Cloud Marketing Hub (CMH), une entreprise du secteur du marketing digital et des solutions technologiques, où nous avons participé au développement de CareerHub, une application web destinée à la gestion des promotions de formation et au suivi des candidats.",
    )
    add_heading(doc, "1.3 Objectifs pédagogiques et professionnels", 2)
    add_paragraph(doc, "Objectifs pédagogiques :", bold=True)
    add_bullets(
        doc,
        [
            "Mettre en œuvre une architecture full stack (React + Laravel) ;",
            "Concevoir et implémenter une API REST sécurisée ;",
            "Modéliser et exploiter une base de données relationnelle ;",
            "Appliquer les principes de l'architecture MVC ;",
            "Travailler en équipe sur un projet versionné avec Git.",
        ],
    )
    add_paragraph(doc, "Objectifs professionnels :", bold=True)
    add_bullets(
        doc,
        [
            "Comprendre les besoins métier d'une entreprise de formation / recrutement ;",
            "Livrer une application fonctionnelle, testée et documentée ;",
            "Adopter les standards de qualité et de sécurité en entreprise ;",
            "Communiquer efficacement avec l'équipe et le tuteur professionnel.",
        ],
    )
    add_heading(doc, "1.4 Attentes personnelles", 2)
    add_paragraph(
        doc,
        "Mon attente principale était de mettre en pratique mes compétences en développement web full stack sur un projet concret, avec des contraintes réelles de délai, de qualité et de collaboration. Je souhaitais notamment maîtriser le développement d'une SPA avec React, comprendre l'intégration front-end / back-end via une API REST, approfondir la sécurité (authentification, validation, protection des données) et contribuer à un produit utile pour l'entreprise d'accueil.",
    )
    doc.add_page_break()


def build_entreprise(doc: Document) -> None:
    add_heading(doc, "2. PRÉSENTATION DE L'ENTREPRISE", 1)
    add_table(
        doc,
        ["Élément", "Information"],
        [
            ["Raison sociale", "Cloud Marketing Hub (CMH)"],
            ["Adresse", "...."],
            ["Date de création", "...."],
            ["Site web", "...."],
            ["Secteur d'activité", "Marketing digital, solutions technologiques, formation"],
        ],
        [Cm(5), Cm(11)],
    )
    add_paragraph(doc, "")
    add_heading(doc, "2.1 Secteur d'activité", 2)
    add_paragraph(
        doc,
        "CMH évolue dans le domaine du marketing digital et des solutions numériques. L'entreprise accompagne ses clients dans la transformation digitale et propose des services liés à la formation, au recrutement et à la gestion des talents. Le projet CareerHub s'inscrit directement dans cette activité : il digitalise le suivi des promotions de formation et l'évaluation des candidats.",
    )
    add_heading(doc, "2.2 Taille et organisation", 2)
    add_paragraph(doc, "L'entreprise compte environ .... collaborateurs, répartis entre la direction, le service développement/IT, le marketing, la formation/RH et ....")
    add_heading(doc, "2.3 Services proposés", 2)
    add_bullets(
        doc,
        [
            "Conception de stratégies marketing digital ;",
            "Développement d'applications web et mobiles ;",
            "Gestion et suivi de programmes de formation ;",
            "Accompagnement des candidats et évaluation des compétences ;",
            "....",
        ],
    )
    doc.add_page_break()


def build_service(doc: Document) -> None:
    add_heading(doc, "3. PRÉSENTATION DU SERVICE D'ACCUEIL", 1)
    add_heading(doc, "3.1 Département intégré", 2)
    add_paragraph(doc, "Nous avons été intégrées au service de développement informatique (....), chargé de la conception et de la maintenance des outils internes et clients de CMH.")
    add_heading(doc, "3.2 Composition de l'équipe", 2)
    add_table(
        doc,
        ["Rôle", "Nom / fonction"],
        [
            ["Tuteur professionnel", "...."],
            ["Développeur(s) senior(s)", "...."],
            ["Équipe de stage", ", ".join(TEAM)],
        ],
        [Cm(5), Cm(11)],
    )
    add_heading(doc, "3.3 Outils de travail", 2)
    add_bullets(doc, ["Visual Studio Code", "GitHub (monorepo cmh_career_hub)", "Postman", ".... (Slack, Teams, Trello, Jira...)"])
    add_heading(doc, "3.4 Méthodologie", 2)
    add_paragraph(doc, "Le projet a été mené selon une approche Agile inspirée de Scrum : sprints, backlog, points d'avancement, livraisons incrémentales et revues de code.")
    doc.add_page_break()


def build_mission_block(doc: Document, title: str, objectives: list[str], tech: str, steps: list[str], results: list[str]) -> None:
    add_heading(doc, title, 2)
    add_paragraph(doc, "Objectifs :", bold=True)
    add_bullets(doc, objectives)
    add_paragraph(doc, f"Technologies utilisées : {tech}")
    add_paragraph(doc, "Étapes de réalisation :", bold=True)
    add_bullets(doc, steps)
    add_paragraph(doc, "Résultats obtenus :", bold=True)
    add_bullets(doc, results)


def build_missions(doc: Document) -> None:
    add_heading(doc, "4. MISSIONS ET TÂCHES RÉALISÉES", 1)
    add_paragraph(
        doc,
        "Le projet principal confié à notre équipe est CMH Career Hub — une plateforme web full stack de gestion des promotions de formation et du parcours des candidats.",
    )
    build_mission_block(
        doc,
        "Mission 1 : Analyse des besoins et conception de l'architecture",
        [
            "Comprendre le processus métier (promotions, candidats, modules, notes, compétences) ;",
            "Définir l'architecture technique front-end / back-end ;",
            "Modéliser la base de données.",
        ],
        "UML, MySQL, papier/....",
        [
            "Entretiens avec le tuteur et analyse du cahier des charges ;",
            "Identification des entités : Promotion, Candidate, Module, ModuleGrade, CandidateSkill, User ;",
            "Architecture monorepo : frontend/ (React) et backend/ (Laravel 11) ;",
            "Définition des endpoints API REST (/api/v1/...).",
        ],
        ["Schéma de base de données validé ;", "Architecture MVC ;", "Plan de développement par sprints."],
    )
    build_mission_block(
        doc,
        "Mission 2 : Développement du back-end (API Laravel)",
        [
            "Créer une API REST sécurisée ;",
            "Implémenter la logique métier (moyennes, catégories, statistiques).",
        ],
        "PHP 8.2, Laravel 11, MySQL, Laravel Sanctum, Eloquent ORM",
        [
            "Initialisation Laravel, configuration .env, CORS, base de données ;",
            "Migrations et modèles Eloquent ;",
            "Contrôleurs : Auth, Candidate, Promotion, Dashboard, Reminders, AI ;",
            "Middlewares : auth:sanctum, rate limiting, en-têtes de sécurité ;",
            "Services métier : CandidateQueryService, PromotionStatsService, AiContextService.",
        ],
        ["API fonctionnelle sur http://127.0.0.1:8000/api/v1/ ;", "Endpoints testés avec Postman ;", "Validation et règles métier (StrongPassword)."],
    )
    build_mission_block(
        doc,
        "Mission 3 : Système d'authentification sécurisé",
        [
            "Protéger l'accès à l'application ;",
            "Gérer les sessions via tokens API.",
        ],
        "Laravel Sanctum, React, Zustand",
        [
            "Endpoint POST /auth/login ;",
            "Génération de token Bearer Sanctum ;",
            "Protection des routes par middleware auth:sanctum ;",
            "Page de connexion React ;",
            "Limitation des tentatives (throttle:login).",
        ],
        ["Authentification front ↔ back opérationnelle ;", "Déconnexion et profil fonctionnels ;", "Routes sensibles sécurisées."],
    )
    build_mission_block(
        doc,
        "Mission 4 : Développement du front-end (React + TypeScript)",
        [
            "Construire une interface moderne, responsive et intuitive ;",
            "Consommer l'API REST de manière fiable.",
        ],
        "React 19, TypeScript, TanStack Router/Query, Tailwind CSS 4, Radix UI, Recharts, Zod",
        [
            "Routes : Dashboard, Promotions, Candidats, Rapports, AI Advisor, Paramètres, Login ;",
            "Composants UI réutilisables ;",
            "Services API et gestion d'état Zustand ;",
            "Thème clair/sombre et i18n FR/EN pour les exports.",
        ],
        ["SPA sur http://127.0.0.1:5173 ;", "Navigation fluide ;", "Interface conforme à l'identité CareerHub / CMH."],
    )
    build_mission_block(
        doc,
        "Mission 5 : Gestion des promotions et candidats",
        [
            "Créer et suivre les promotions de formation ;",
            "Gérer le cycle de vie complet d'un candidat.",
        ],
        "Laravel Eloquent, Soft Deletes, React Hook Form",
        [
            "CRUD promotions (code, dates, statut) ;",
            "CRUD candidats (identité, contact, études, photo, état) ;",
            "Notes par module et calcul automatique de la moyenne ;",
            "Compétences par catégorie ;",
            "Archivage / restauration.",
        ],
        ["Gestion complète promotions/candidats ;", "Calcul automatique des catégories ;", "ModuleGradeObserver pour recalcul des moyennes."],
    )
    build_mission_block(
        doc,
        "Mission 6 : Dashboard et statistiques",
        ["Offrir une vue d'ensemble aux responsables formation ;", "Visualiser les KPIs et tendances."],
        "Recharts, API GET /dashboard/stats, React Query",
        [
            "Statistiques : actifs, diplômés, terminés, taux de réussite ;",
            "Graphiques : genre, âge, éducation, scores par promotion ;",
            "Cartes KPI et rappels contextuels (RemindersService).",
        ],
        ["Tableau de bord interactif ;", "Aide à la décision pour le suivi pédagogique."],
    )
    build_mission_block(
        doc,
        "Mission 7 : Rapports et exports",
        ["Exporter les données candidat et promotion ;", "Supporter plusieurs formats et langues."],
        "jsPDF, jspdf-autotable, xlsx, CandidateExportService",
        [
            "Page Reports & Exports ;",
            "Formats PDF, Excel, HTML ;",
            "Localisation FR/EN ;",
            "Endpoint GET /candidates/export.",
        ],
        ["Exports téléchargeables ;", "Rapports utilisables en contexte professionnel."],
    )
    build_mission_block(
        doc,
        "Mission 8 : Conseiller IA (AI Advisor)",
        ["Intégrer une assistance intelligente basée sur les données ;", "Fournir des analyses contextuelles."],
        "Laravel AiController/AiContextService, React ai-advisor",
        [
            "Construction du contexte analytics ;",
            "Endpoint POST /ai/chat protégé (throttle:ai) ;",
            "Interface de chat ;",
            "Modèle IA : ....",
        ],
        ["Module d'assistance opérationnel ;", "Enrichissement de la valeur métier de CareerHub."],
    )
    build_mission_block(
        doc,
        "Mission 9 : Tests, déploiement et documentation",
        ["Vérifier le bon fonctionnement ;", "Documenter l'installation et l'utilisation."],
        "PHPUnit, Postman, README, Git",
        [
            "Tests manuels des parcours utilisateur ;",
            "Script smoke test base de données ;",
            "README monorepo ;",
            "Environnement de déploiement : ....",
        ],
        ["Application stable en local ;", "Documentation pour reprise du projet."],
    )
    doc.add_page_break()


def build_technologies(doc: Document) -> None:
    add_heading(doc, "5. OUTILS, TECHNOLOGIES ET MÉTHODES UTILISÉS", 1)
    add_table(
        doc,
        ["Catégorie", "Technologies"],
        [
            ["Langages", "HTML, CSS, JavaScript, TypeScript, PHP, SQL"],
            ["Frameworks", "React 19, Laravel 11, TanStack Router/Query, Tailwind CSS 4"],
            ["Base de données", "MySQL"],
            ["Outils", "GitHub, VS Code, Postman, Bun, Composer, Vite"],
            ["Méthodes", "MVC, REST API, Agile/Scrum, Soft Delete, Observer Pattern"],
        ],
        [Cm(4), Cm(12)],
    )
    doc.add_page_break()


def build_competences(doc: Document) -> None:
    add_heading(doc, "6. COMPÉTENCES ACQUISES", 1)
    add_heading(doc, "6.1 Compétences techniques", 2)
    add_bullets(
        doc,
        [
            "Développement full stack React + Laravel ;",
            "Conception et implémentation d'une API REST sécurisée ;",
            "Modélisation base de données relationnelle ;",
            "Authentification par tokens (Sanctum) ;",
            "Dashboards avec graphiques interactifs ;",
            "Export multi-formats (PDF, Excel, HTML) ;",
            "TypeScript, CORS, middlewares de sécurité ;",
            "Travail en monorepo.",
        ],
    )
    add_heading(doc, "6.2 Compétences transversales", 2)
    add_table(
        doc,
        ["Compétence", "Description"],
        [
            ["Travail en équipe", "Répartition des tâches entre 4 stagiaires, revues mutuelles, Git"],
            ["Communication", "Points réguliers avec le tuteur, clarification des besoins métier"],
            ["Gestion du temps", "Respect des jalons, priorisation des fonctionnalités critiques"],
            ["Résolution de problèmes", "Debug API, CORS, cohérence front/back, documentation"],
            ["Autonomie", "Recherche de solutions, lecture de documentation officielle"],
            ["Professionnalisme", "Code structuré, conventions, livrables documentés"],
        ],
        [Cm(4.5), Cm(11.5)],
    )
    doc.add_page_break()


def build_difficultes(doc: Document) -> None:
    add_heading(doc, "7. DIFFICULTÉS RENCONTRÉES ET SOLUTIONS APPORTÉES", 1)
    difficulties = [
        (
            "Intégration front-end / back-end (CORS et authentification)",
            "Erreurs CORS et échecs d'authentification lors des premiers appels API.",
            "Configuration de config/cors.php, envoi du token Bearer, tests Postman avant intégration front.",
        ),
        (
            "Synchronisation des moyennes candidats",
            "Recalcul automatique de la moyenne globale après modification d'une note.",
            "ModuleGradeObserver Laravel + CandidateScoreService centralisé.",
        ),
        (
            "Exports PDF/Excel multilingues",
            "Cohérence des libellés FR/EN selon le format.",
            "Fichiers export-labels.ts et ExportLocale.php, factorisation des exports.",
        ),
        (
            "Gestion d'état complexe côté front",
            "Données partagées entre plusieurs pages.",
            "TanStack Query pour le serveur + Zustand pour l'état global.",
        ),
        (
            "Répartition du travail en équipe de 4",
            "Risque de conflits Git et doublons.",
            "Branches par fonctionnalité, communication quotidienne, conventions de nommage.",
        ),
    ]
    for title, problem, solution in difficulties:
        add_paragraph(doc, title, bold=True)
        add_paragraph(doc, f"Problème : {problem}")
        add_paragraph(doc, f"Solution : {solution}")
        add_paragraph(doc, "")
    doc.add_page_break()


def build_bilan(doc: Document) -> None:
    add_heading(doc, "8. BILAN PERSONNEL", 1)
    sections = [
        (
            "8.1 Apports du stage",
            "Ce stage m'a permis de vivre une expérience professionnelle complète, de l'analyse des besoins à la livraison d'une application web fonctionnelle. Le projet CareerHub m'a donné une vision concrète du métier de développeuse web full stack.",
        ),
        (
            "8.2 Forces développées",
            "Capacité à développer une application de bout en bout, rigueur en sécurisation, aptitude au travail en équipe sur Git, autonomie dans la recherche de solutions, sens du détail (UX, exports, gestion des erreurs).",
        ),
        (
            "8.3 Axes d'amélioration",
            "Approfondir les tests automatisés (PHPUnit, tests front-end), améliorer les compétences DevOps (CI/CD, déploiement cloud), renforcer la documentation API (OpenAPI/Swagger), ....",
        ),
        (
            "8.4 Vision professionnelle future",
            "À l'issue de ce stage, je souhaite .... (poursuivre en tant que développeuse web full stack / intégrer CMH / ....). Cette expérience confirme mon intérêt pour le développement d'applications web modernes.",
        ),
    ]
    for title, text in sections:
        add_heading(doc, title, 2)
        add_paragraph(doc, text)
    doc.add_page_break()


def build_conclusion(doc: Document) -> None:
    add_heading(doc, "9. CONCLUSION GÉNÉRALE", 1)
    add_paragraph(
        doc,
        "Ce stage de fin de formation au sein de Cloud Marketing Hub (CMH) a constitué une étape enrichissante dans mon parcours à la filière Développement Digital — Web Full Stack au CMCTTA.",
    )
    add_paragraph(
        doc,
        "En équipe avec AYA AOULAD AHMIDOU, ZAHIRA KTAIB et IHSSAN YAKOUBI, nous avons conçu et développé CMH Career Hub, une plateforme web full stack permettant de gérer les promotions, le suivi des candidats, les statistiques, les exports et l'assistance par IA.",
    )
    add_paragraph(doc, "Objectifs atteints :", bold=True)
    add_bullets(
        doc,
        [
            "Application des compétences full stack sur un projet réel ;",
            "Livraison de modules fonctionnels (auth, CRUD, dashboard, rapports, AI) ;",
            "Travail collaboratif professionnel avec Git et méthode Agile ;",
            "Compréhension du lien entre développement technique et besoins métier.",
        ],
    )
    add_paragraph(
        doc,
        "Ce stage a comblé le fossé entre la formation théorique et la pratique professionnelle. Le projet CareerHub peut évoluer vers .... (déploiement production, application mobile, module RH avancé). Je remercie CMH, mon tuteur professionnel, l'équipe pédagogique du CMCTTA et mes collègues de stage pour cette expérience formatrice.",
    )
    doc.add_page_break()


def build_uml_annex(doc: Document) -> None:
    add_heading(doc, "10. ANNEXES", 1)
    add_heading(doc, "Annexe A — Diagramme de classes UML (modèle de données)", 2)
    add_paragraph(
        doc,
        "Le diagramme ci-dessous représente les entités principales de CareerHub et leurs relations, conformément aux migrations Laravel du projet.",
    )

    class_rows = [
        ["User", "id, full_name, email, password, role, timestamps, softDeletes"],
        ["Promotion", "id, promo_code, name, start_date, end_date, status, timestamps, softDeletes"],
        ["Module", "id, promotion_id (FK), name, module_date_debut, module_date_fin, status, module_order"],
        ["Candidate", "id, promotion_id (FK), first_name, last_name, email, phone, recruitment_date, education_level, diploma_specialty, diploma_average, state, overall_avg, category, timestamps, softDeletes"],
    ]
    class_rows.extend([
        ["ModuleGrade", "id, candidate_id (FK), module_id (FK), score (0-20), timestamps — unique(candidate_id, module_id)"],
        ["CandidateSkill", "id, candidate_id (FK), category, skill_name, score (0-5), timestamps — unique(candidate_id, skill_name)"],
    ])
    add_table(doc, ["Classe / Entité", "Attributs principaux"], class_rows, [Cm(4), Cm(12)])

    add_paragraph(doc, "")
    add_paragraph(doc, "Relations UML :", bold=True)
    add_bullets(
        doc,
        [
            "Promotion 1 —— * Module (une promotion contient plusieurs modules)",
            "Promotion 1 —— * Candidate (une promotion contient plusieurs candidats)",
            "Candidate 1 —— * ModuleGrade (un candidat a plusieurs notes par module)",
            "Module 1 —— * ModuleGrade (un module reçoit plusieurs notes)",
            "Candidate 1 —— * CandidateSkill (un candidat possède plusieurs compétences évaluées)",
            "User : entité indépendante pour l'authentification administrateur (Sanctum tokens)",
        ],
    )

    add_paragraph(doc, "")
    add_paragraph(doc, "Représentation textuelle du diagramme de classes :", bold=True)
    uml_text = """
┌─────────────────┐       ┌─────────────────┐       ┌─────────────────┐
│    Promotion    │1     *│     Module      │1     *│  ModuleGrade    │
│─────────────────│◄──────│─────────────────│◄──────│─────────────────│
│ id              │       │ id              │       │ id              │
│ promo_code      │       │ promotion_id    │       │ candidate_id    │
│ name            │       │ name            │       │ module_id       │
│ start_date      │       │ module_date_*   │       │ score           │
│ end_date        │       │ status          │       └────────▲────────┘
│ status          │       │ module_order    │                │
└────────┬────────┘       └─────────────────┘                │
         │1                                                  │*
         │*                                                   │
┌────────▼────────┐                                          │
│   Candidate     │──────────────────────────────────────────┘
│─────────────────│1     *
│ id              │◄──────┌─────────────────┐
│ promotion_id    │       │ CandidateSkill  │
│ first_name      │       │─────────────────│
│ last_name       │       │ candidate_id    │
│ email, phone    │       │ category        │
│ state           │       │ skill_name      │
│ overall_avg     │       │ score           │
│ category        │       └─────────────────┘
└─────────────────┘

┌─────────────────┐
│      User       │  (authentification API — Sanctum)
│─────────────────│
│ id, full_name   │
│ email, password │
│ role            │
└─────────────────┘
"""
    p = doc.add_paragraph()
    run = p.add_run(uml_text.strip())
    run.font.name = "Courier New"
    run.font.size = Pt(8)

    add_heading(doc, "Annexe B — Diagramme de cas d'utilisation", 2)
    add_table(
        doc,
        ["Acteur", "Cas d'utilisation"],
        [
            ["Administrateur", "Se connecter / Se déconnecter / Gérer son profil / Changer le mot de passe"],
            ["Administrateur", "Consulter le dashboard (KPIs, graphiques, rappels)"],
            ["Administrateur", "CRUD Promotions (créer, modifier, archiver, restaurer)"],
            ["Administrateur", "CRUD Candidats (ajouter, modifier, supprimer, restaurer)"],
            ["Administrateur", "Saisir notes par module et compétences candidat"],
            ["Administrateur", "Exporter rapports candidat/promotion (PDF, Excel, HTML — FR/EN)"],
            ["Administrateur", "Utiliser le conseiller IA (AI Advisor)"],
            ["Administrateur", "Configurer les paramètres système"],
            ["Système", "Recalculer automatiquement overall_avg et category (Observer)"],
            ["Système", "Limiter les tentatives de connexion et requêtes IA (throttle)"],
        ],
        [Cm(4), Cm(12)],
    )

    add_paragraph(doc, "")
    add_paragraph(doc, "Diagramme de cas d'utilisation (vue simplifiée) :", bold=True)
    use_case_text = """
                    ┌──────────────────────────────────────────────┐
                    │           CMH Career Hub (Système)           │
                    │                                              │
   ┌──────────┐     │   (Authentification) Se connecter            │
   │          │────►│   (Authentification) Gérer profil            │
   │  Admin   │     │   (Dashboard) Consulter statistiques       │
   │          │────►│   (Promotions) Gérer promotions              │
   │          │     │   (Candidats) Gérer candidats                │
   │          │────►│   (Notes) Saisir notes et compétences        │
   │          │     │   (Exports) Générer rapports PDF/Excel/HTML  │
   │          │────►│   (IA) Interroger le conseiller AI Advisor   │
   └──────────┘     │   (Paramètres) Configurer l'application      │
                    └──────────────────────────────────────────────┘
"""
    p = doc.add_paragraph()
    run = p.add_run(use_case_text.strip())
    run.font.name = "Courier New"
    run.font.size = Pt(8)

    add_heading(doc, "Annexe C — Diagramme de séquence : Authentification", 2)
    seq_text = """
Stagiaire/Admin    Frontend (React)       Backend (Laravel)        MySQL
      │                   │                      │                    │
      │── email/password ─►│                      │                    │
      │                   │── POST /auth/login ─►│                    │
      │                   │                      │── SELECT user ────►│
      │                   │                      │◄── user row ───────│
      │                   │                      │── verify password  │
      │                   │                      │── create Sanctum token
      │                   │◄── {user, token} ────│                    │
      │                   │── store token        │                    │
      │◄── redirect dash ─│                      │                    │
      │                   │── GET /dashboard ───►│ (Bearer token)     │
      │                   │◄── stats JSON ───────│                    │
"""
    p = doc.add_paragraph()
    run = p.add_run(seq_text.strip())
    run.font.name = "Courier New"
    run.font.size = Pt(9)

    add_heading(doc, "Annexe D — Architecture technique", 2)
    add_table(
        doc,
        ["Couche", "Technologie", "Rôle"],
        [
            ["Présentation", "React 19 + TypeScript + Tailwind", "Interface utilisateur SPA"],
            ["Routing / État", "TanStack Router, Query, Zustand", "Navigation et cache API"],
            ["API", "Laravel 11 REST /api/v1", "Logique métier et endpoints JSON"],
            ["Sécurité", "Laravel Sanctum", "Authentification par token Bearer"],
            ["Données", "MySQL + Eloquent ORM", "Persistance relationnelle"],
            ["Outils", "GitHub, VS Code, Postman", "Collaboration et tests"],
        ],
        [Cm(3.5), Cm(5), Cm(7.5)],
    )

    doc.add_page_break()


def build_planning_annex(doc: Document) -> None:
    add_heading(doc, "Annexe E — Planning de stage et répartition des tâches (équipe de 4)", 2)
    add_paragraph(
        doc,
        "Période du stage : du .... au .... — Durée totale : .... semaines. Méthode : Agile (sprints de .... semaine(s)).",
    )

    add_heading(doc, "E.1 Planning global par phases", 3)
    planning_rows = [
        ["Semaine 1", "Analyse des besoins, maquettes, modélisation BDD, setup monorepo", "Toute l'équipe"],
        ["Semaine 2", "Back-end Laravel : migrations, modèles, API auth (Sanctum)", "AYA EL OUAHABI + IHSSAN YAKOUBI"],
        ["Semaine 3", "Front-end : routing, login, layout, sidebar, thème", "AYA AOULAD AHMIDOU + ZAHIRA KTAIB"],
        ["Semaine 4", "Module Promotions (CRUD, archivage, stats)", "ZAHIRA KTAIB + AYA EL OUAHABI"],
        ["Semaine 5", "Module Candidats (CRUD, notes, compétences, fiche détail)", "AYA AOULAD AHMIDOU + IHSSAN YAKOUBI"],
        ["Semaine 6", "Dashboard (KPIs, graphiques Recharts, rappels)", "AYA AOULAD AHMIDOU + AYA EL OUAHABI"],
        ["Semaine 7", "Exports PDF/Excel/HTML (FR/EN)", "ZAHIRA KTAIB + AYA AOULAD AHMIDOU"],
        ["Semaine 8", "AI Advisor, paramètres, tests, documentation, préparation rapport", "IHSSAN YAKOUBI + toute l'équipe"],
    ]
    add_table(doc, ["Période", "Activités / Livrables", "Responsables"], planning_rows, [Cm(2.5), Cm(9.5), Cm(4)])

    add_heading(doc, "E.2 Répartition détaillée par stagiaire", 3)

    members_tasks = [
        (
            "AYA EL OUAHABI — Coordinatrice technique / Back-end",
            [
                "Analyse des besoins et conception architecture monorepo ;",
                "API Laravel : AuthController, middlewares, StrongPassword ;",
                "Services : CandidateQueryService, sécurité (Sanctum, throttle) ;",
                "Dashboard API (DashboardController, stats) ;",
                "Revues de code et merge GitHub ;",
                "Rédaction du rapport de stage.",
            ],
        ),
        (
            "AYA AOULAD AHMIDOU — Front-end / UX Dashboard & Candidats",
            [
                "Structure React + TanStack Router ;",
                "Page Login et gestion auth côté front (Zustand) ;",
                "Dashboard : KPI cards, graphiques Recharts, filtres ;",
                "Pages candidats (liste, fiche détail, formulaires) ;",
                "Composants UI réutilisables (loading states, badges) ;",
                "Tests manuels parcours utilisateur.",
            ],
        ),
        (
            "ZAHIRA KTAIB — Promotions & Exports",
            [
                "CRUD Promotions (PromotionController + pages React) ;",
                "Archivage / restauration promotions ;",
                "Module Reports & Exports (PDF, Excel, HTML) ;",
                "Internationalisation exports FR/EN ;",
                "PromotionStatsService et export-data ;",
                "Captures d'écran pour annexes du rapport.",
            ],
        ),
        (
            "IHSSAN YAKOUBI — Back-end métier / IA & Qualité",
            [
                "Migrations et modèles (Candidate, Module, ModuleGrade, CandidateSkill) ;",
                "ModuleGradeObserver et CandidateScoreService ;",
                "RemindersService et page rappels ;",
                "AI Advisor (AiController, AiContextService, interface chat) ;",
                "Tests API Postman, script db-smoke-test ;",
                "Documentation README et annexes UML.",
            ],
        ),
    ]
    for title, tasks in members_tasks:
        add_paragraph(doc, title, bold=True)
        add_bullets(doc, tasks)
        add_paragraph(doc, "")

    add_heading(doc, "E.3 Matrice RACI simplifiée", 3)
    add_table(
        doc,
        ["Module / Tâche", "AYA E.O.", "AYA A.A.", "ZAHIRA K.", "IHSSAN Y."],
        [
            ["Architecture & setup", "R/A", "C", "C", "C"],
            ["Authentification API", "R/A", "I", "I", "C"],
            ["Front-end layout & routing", "C", "R/A", "C", "I"],
            ["Module Promotions", "C", "I", "R/A", "C"],
            ["Module Candidats", "C", "R/A", "C", "C"],
            ["Dashboard & stats", "R", "R/A", "C", "C"],
            ["Exports & rapports", "C", "C", "R/A", "I"],
            ["AI Advisor", "C", "I", "I", "R/A"],
            ["Tests & documentation", "A", "C", "C", "R"],
        ],
        [Cm(4.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5)],
    )
    add_paragraph(doc, "Légende : R = Responsable, A = Approbateur, C = Contributeur, I = Informé")

    add_heading(doc, "Annexe F — Liste des captures d'écran à insérer", 3)
    add_bullets(
        doc,
        [
            "Page de connexion CareerHub ;",
            "Dashboard (KPIs et graphiques) ;",
            "Liste et fiche candidat ;",
            "Gestion des promotions ;",
            "Page Reports & Exports ;",
            "AI Advisor ;",
            "Attestation de stage ;",
            "....",
        ],
    )
    doc.add_page_break()


def build_bibliographie(doc: Document) -> None:
    add_heading(doc, "11. BIBLIOGRAPHIE / WEBOGRAPHIE", 1)
    refs = [
        "Laravel Documentation — https://laravel.com/docs/11.x",
        "React Documentation — https://react.dev",
        "PHP Manual — https://www.php.net/manual/fr/",
        "MySQL Documentation — https://dev.mysql.com/doc/",
        "TypeScript Handbook — https://www.typescriptlang.org/docs/",
        "TanStack Router — https://tanstack.com/router",
        "TanStack Query — https://tanstack.com/query",
        "Tailwind CSS — https://tailwindcss.com/docs",
        "Laravel Sanctum — https://laravel.com/docs/11.x/sanctum",
        "MDN Web Docs — https://developer.mozilla.org/fr/",
        "Recharts — https://recharts.org",
        "Zod — https://zod.dev",
        "Stack Overflow — https://stackoverflow.com",
        "GitHub — https://github.com (dépôt cmh_career_hub)",
        "OFPPT — https://www.ofppt.ma",
        "....",
    ]
    for i, ref in enumerate(refs, 1):
        add_paragraph(doc, f"{i}. {ref}")


def validate_docx(path: Path) -> None:
    import zipfile
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(path) as zf:
        if "word/document.xml" not in zf.namelist():
            raise ValueError("document.xml manquant")
        ET.fromstring(zf.read("word/document.xml"))
        for name in zf.namelist():
            if name.startswith("word/footer") and name.endswith(".xml"):
                ET.fromstring(zf.read(name))


def generate() -> Path:
    doc = Document()
    set_document_defaults(doc)
    add_simple_footer(doc)

    build_cover_page(doc)
    build_remerciements(doc)
    build_sommaire(doc)
    build_introduction(doc)
    build_entreprise(doc)
    build_service(doc)
    build_missions(doc)
    build_technologies(doc)
    build_competences(doc)
    build_difficultes(doc)
    build_bilan(doc)
    build_conclusion(doc)
    build_uml_annex(doc)
    build_planning_annex(doc)
    build_bibliographie(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    validate_docx(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = generate()
    print(f"Rapport genere : {path}")
